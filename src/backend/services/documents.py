from io import BytesIO
from pathlib import Path


def extract_document(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in {".txt", ".md", ".markdown", ".prd"}:
        return data.decode("utf-8", errors="replace")
    if ext == ".docx":
        return _docx(data)
    if ext == ".pdf":
        return _pdf(data)
    raise ValueError("仅支持 PDF、Word（.docx）、Markdown、TXT")


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)
